# Code Iteration #39
import csv
import os
import sys
import math

# -----------------------------------------------------------------------------
# 1. Automatic Package Resolver
# -----------------------------------------------------------------------------
try:
    import reportlab
except ImportError:
    print("Notice: 'reportlab' module not found. Installing reportlab for direct PDF generation...", flush=True)
    try:
        import subprocess
        subprocess.run([sys.executable, "-m", "pip", "install", "--user", "reportlab"], check=True)
        import reportlab
        print("Success: 'reportlab' installed successfully!", flush=True)
    except Exception as e:
        print(f"Warning: Could not auto-install reportlab via pip ({e}).", flush=True)

from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

# Maximized vertical capacities to extend tables to the bottom
MAX_ROWS_PER_PDF_PAGE = 38
MAX_ROWS_PER_PPTX_SLIDE = 22

def set_pptx_cell_border(cell, color="000000", width="12700"):
    """Applies explicit OpenXML stroke lines to PowerPoint table cells."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcBorders'):
            tcPr.remove(child)
            
    tcBorders = parse_xml(f'''
        <a:tcBorders {nsdecls("a")}>
            <a:lnL w="{width}" cmpd="s"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:lnL>
            <a:lnR w="{width}" cmpd="s"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:lnR>
            <a:lnT w="{width}" cmpd="s"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:lnT>
            <a:lnB w="{width}" cmpd="s"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:lnB>
        </a:tcBorders>
    ''')
    tcPr.append(tcBorders)

def evaluate_pass_fail(val_str, spec_str):
    """Evaluates numeric values against spec strings strictly."""
    if not spec_str or not spec_str.strip():
        return None
        
    try:
        val = float(val_str)
        spec = spec_str.strip()
        
        if ".." in spec or " to " in spec:
            parts = spec.replace(" to ", "..").split("..")
            low, high = float(parts[0].strip()), float(parts[1].strip())
            return low <= val <= high
        elif spec.startswith(">="):
            return val >= float(spec.replace(">=", "").strip())
        elif spec.startswith(">"):
            return val > float(spec.replace(">", "").strip())
        elif spec.startswith("<="):
            return val <= float(spec.replace("<=", "").strip())
        elif spec.startswith("<"):
            return val < float(spec.replace("<", "").strip())
    except ValueError:
        pass
        
    return None

def build_reports(csv_path, pptx_output_path):
    print("\n" + "=" * 70, flush=True)
    print("WARNING: Make sure to SAVE your latest .csv export from Cadence Maestro first!", flush=True)
    print("=" * 70 + "\n", flush=True)

    full_csv_path = os.path.expanduser(csv_path)
    if not os.path.exists(full_csv_path):
        print(f"ERROR: Could not find CSV file at '{full_csv_path}'", flush=True)
        return

    headers = []
    filtered_data = []

    # Read and parse CSV
    with open(full_csv_path, 'r') as f:
        reader = csv.reader(f)
        header_found = False
        raw_headers = []
        keep_indices = []
        
        for row in reader:
            if not row:
                continue
            
            if len(row) > 1 and row[0].strip() == "Test" and row[1].strip() == "Output":
                raw_headers = [col.strip() for col in row if col.strip() != ""]
                keep_indices = [i for i, h in enumerate(raw_headers) if h not in ["Test", "Weight", "Pass/Fail"]]
                headers = [raw_headers[i] for i in keep_indices]
                header_found = True
                continue
            
            if header_found and len(row) > 1 and row[1].strip() != "":
                full_row = [col.strip() for col in row[:len(raw_headers)]]
                corner_vals = full_row[7:] if len(full_row) > 7 else full_row[2:]
                
                has_valid_data = False
                has_eval_error = False
                
                for cv in corner_vals:
                    val_lower = cv.lower().strip()
                    if "eval" in val_lower or "err" in val_lower:
                        has_eval_error = True
                        break
                    if val_lower != "":
                        has_valid_data = True

                if has_valid_data and not has_eval_error:
                    clean_row = [full_row[i] for i in keep_indices if i < len(full_row)]
                    filtered_data.append(clean_row)

    if not headers or not filtered_data:
        print("ERROR: No valid numerical data remaining after filtering.", flush=True)
        return

    spec_col_idx = headers.index("Spec") if "Spec" in headers else -1
    total_data_rows = len(filtered_data)

    # -----------------------------------------------------------------------------
    # 2. Build PDF Document using ReportLab (Col 1 = 135pt, 38 Rows/Page)
    # -----------------------------------------------------------------------------
    pdf_output_path = pptx_output_path.rsplit('.', 1)[0] + ".pdf"
    full_pdf_path = os.path.expanduser(pdf_output_path)
    
    try:
        doc = SimpleDocTemplate(
            full_pdf_path,
            pagesize=landscape(letter),
            leftMargin=20, rightMargin=20, topMargin=15, bottomMargin=15
        )
        
        styles = getSampleStyleSheet()
        header_paragraph_style = ParagraphStyle(
            'HeaderStyle', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=8.5, leading=10,
            textColor=colors.white, alignment=0
        )
        cell_paragraph_style = ParagraphStyle(
            'CellStyle', parent=styles['Normal'],
            fontName='Helvetica', fontSize=8.0, leading=9.5,
            textColor=colors.black, alignment=0
        )
        
        pdf_story = []
        num_pdf_pages = math.ceil(total_data_rows / MAX_ROWS_PER_PDF_PAGE)

        for page_idx in range(num_pdf_pages):
            start_row = page_idx * MAX_ROWS_PER_PDF_PAGE
            end_row = min(start_row + MAX_ROWS_PER_PDF_PAGE, total_data_rows)
            page_rows = filtered_data[start_row:end_row]

            table_data = [[Paragraph(h_text, header_paragraph_style) for h_text in headers]]
            table_style_cmd = [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3C78')),
                ('GRID', (0, 0), (-1, -1), 0.8, colors.black),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 1.8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 1.8),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ]

            for r_idx, row_values in enumerate(page_rows, start=1):
                spec_val = row_values[spec_col_idx] if (spec_col_idx != -1 and spec_col_idx < len(row_values)) else ""
                row_cells = []
                
                for c_idx in range(len(headers)):
                    val_text = row_values[c_idx] if c_idx < len(row_values) else ""
                    row_cells.append(Paragraph(val_text, cell_paragraph_style))

                    is_corner_col = (c_idx > spec_col_idx) if spec_col_idx != -1 else True
                    is_pass = evaluate_pass_fail(val_text, spec_val) if is_corner_col else None
                    
                    if is_pass is True:
                        table_style_cmd.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), colors.HexColor('#34FC32')))
                    elif is_pass is False:
                        table_style_cmd.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), colors.HexColor('#FA2F2F')))

                table_data.append(row_cells)

            # Usable width = 752pt. Set Col 0 ("Output") to 135pt; split remainder evenly across other cols.
            if len(headers) > 1:
                rem_width = (752.0 - 135.0) / (len(headers) - 1)
                col_widths = [135.0] + [rem_width] * (len(headers) - 1)
            else:
                col_widths = [752.0]

            pdf_table = Table(table_data, colWidths=col_widths)
            pdf_table.setStyle(TableStyle(table_style_cmd))
            
            pdf_story.append(pdf_table)
            if page_idx < num_pdf_pages - 1:
                pdf_story.append(PageBreak())

        doc.build(pdf_story)
        print(f"SUCCESS: Exported PDF report directly via ReportLab to '{full_pdf_path}'", flush=True)
    except Exception as e:
        print(f"NOTICE: Direct ReportLab PDF export encountered an issue ({e}).", flush=True)

    # -----------------------------------------------------------------------------
    # 3. Build PowerPoint Deck (.pptx) with Balanced Col 1 Width & 22 Rows/Slide
    # -----------------------------------------------------------------------------
    prs = Presentation()
    num_cols = len(headers)
    num_pptx_slides = math.ceil(total_data_rows / MAX_ROWS_PER_PPTX_SLIDE)

    for slide_idx in range(num_pptx_slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        start_row = slide_idx * MAX_ROWS_PER_PPTX_SLIDE
        end_row = min(start_row + MAX_ROWS_PER_PPTX_SLIDE, total_data_rows)
        slide_rows = filtered_data[start_row:end_row]
        rows_in_table = len(slide_rows) + 1

        left = Inches(0.4)
        top = Inches(0.3)
        total_width = Inches(9.2)
        height = Inches(0.28 * rows_in_table)

        table_shape = slide.shapes.add_table(rows_in_table, num_cols, left, top, total_width, height)
        table = table_shape.table

        # Balanced Column Widths: Col 0 gets 1.8 inches; split remainder across corner cols
        if num_cols > 1:
            col0_width_int = int(Inches(1.8))
            rem_col_width_int = int((Inches(9.2) - Inches(1.8)) / (num_cols - 1))
            
            table.columns[0].width = col0_width_int
            for c_i in range(1, num_cols):
                table.columns[c_i].width = rem_col_width_int

        # Strips default PowerPoint table theme override
        tblPr = table._tbl.tblPr
        tblPr.set('firstRow', '0')
        tblPr.set('bandRow', '0')
        style_id = tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
        if style_id is not None:
            tblPr.remove(style_id)

        # Header Row Styling
        for col_idx, h_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = h_text
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(30, 60, 120)
            
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.font.size = Pt(9.5)

            set_pptx_cell_border(cell, color="000000", width="19050")

        # Data Rows Styling
        for r_idx, row_values in enumerate(slide_rows, start=1):
            spec_val = row_values[spec_col_idx] if (spec_col_idx != -1 and spec_col_idx < len(row_values)) else ""

            for c_idx in range(min(num_cols, len(row_values))):
                cell = table.cell(r_idx, c_idx)
                val_text = row_values[c_idx]
                cell.text = val_text

                is_corner_col = (c_idx > spec_col_idx) if spec_col_idx != -1 else True
                is_pass = evaluate_pass_fail(val_text, spec_val) if is_corner_col else None
                
                if is_pass is True:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0x34, 0xFC, 0x32)
                elif is_pass is False:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xFA, 0x2F, 0x2F)
                else:
                    cell.fill.background()

                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(8.5)
                    p.font.color.rgb = RGBColor(0, 0, 0)

                set_pptx_cell_border(cell, color="000000", width="12700")

    out_pptx = os.path.expanduser(pptx_output_path)
    prs.save(out_pptx)
    print(f"SUCCESS: Created PowerPoint deck at '{out_pptx}'", flush=True)

if __name__ == "__main__":
    run_num = sys.argv[1] if len(sys.argv) > 1 else "430"
    input_csv = f"~/Documents/Interactive.{run_num}.csv"
    output_presentation = f"~/Documents/Simulation_Summary_{run_num}.pptx"
    
    build_reports(input_csv, output_presentation)
